"""Downloadable PDF/DOCX reports shared by the new web result screens."""
from io import BytesIO
import json
from xml.sax.saxutils import escape


def report_bytes(title, records, format):
    sections = []
    for index, record in enumerate(records, 1):
        sections.append((f"{index}. {record.get('question', 'Evaluation')}", json.dumps(record, ensure_ascii=False, indent=2)))
    output = BytesIO()
    if format == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        styles = getSampleStyleSheet()
        story = [Paragraph(escape(title), styles['Title']), Spacer(1, 16)]
        for heading, content in sections:
            story.append(Paragraph(escape(heading), styles['Heading2']))
            for line in content.splitlines():
                story.append(Paragraph(escape(line) or ' ', styles['BodyText']))
            story.append(Spacer(1, 14))
        SimpleDocTemplate(output).build(story)
    elif format == 'docx':
        from docx import Document
        document = Document()
        document.add_heading(title, 0)
        for heading, content in sections:
            document.add_heading(heading, 1)
            document.add_paragraph(content)
        document.save(output)
    else:
        raise ValueError('Choose PDF or DOCX')
    return output.getvalue()
